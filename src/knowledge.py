from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import chromadb
from chromadb.config import Settings
from docx import Document as WordDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader


def load_document(name: str, data: bytes) -> list[Document]:
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Each document must be 10 MB or smaller.')
    suffix, source = Path(name).suffix.lower(), Path(name).name
    if suffix == '.pdf':
        pdf = PdfReader(BytesIO(data))
        if len(pdf.pages) > 100:
            raise ValueError('Use a PDF with no more than 100 pages.')
        docs = [Document(page_content=page.extract_text() or '', metadata={'source':source,'page':i+1})
                for i,page in enumerate(pdf.pages)]
    elif suffix == '.docx':
        word = WordDocument(BytesIO(data))
        text = '\n'.join([p.text for p in word.paragraphs] +
                         [' | '.join(c.text for c in r.cells) for t in word.tables for r in t.rows])
        docs = [Document(page_content=text,metadata={'source':source})]
    elif suffix in {'.md','.txt'}:
        docs = [Document(page_content=data.decode('utf-8-sig'),metadata={'source':source})]
    else:
        raise ValueError('Use PDF, DOCX, Markdown or TXT.')
    docs = [d for d in docs if d.page_content.strip()]
    if not docs:
        raise ValueError('No readable text found. Scan-based PDFs need OCR first.')
    if sum(len(d.page_content) for d in docs) > 200000:
        raise ValueError('Each document must contain at most 200,000 characters.')
    return docs


class SupportKnowledge:
    def __init__(self, documents: list[Document], key: str, model: str = 'gemini-embedding-001', embeddings=None):
        if not documents:
            raise ValueError('Add product documentation before indexing.')
        if sum(len(d.page_content) for d in documents) > 400000:
            raise ValueError('The combined knowledge base may contain at most 400,000 characters.')
        if not key and embeddings is None:
            raise ValueError('Gemini credentials are required for semantic embeddings.')
        self.chunks = RecursiveCharacterTextSplitter(chunk_size=850,chunk_overlap=120).split_documents(documents)
        for index, doc in enumerate(self.chunks):
            identity = str(doc.metadata) + doc.page_content
            doc.metadata.update(id='D-'+hashlib.sha256(identity.encode()).hexdigest()[:8],chunk=index+1)
        embed = embeddings or GoogleGenerativeAIEmbeddings(model=model,api_key=key,vertexai=False)
        self.client = chromadb.EphemeralClient(settings=Settings(anonymized_telemetry=False))
        self.store = Chroma(collection_name='support_'+uuid4().hex,client=self.client,
                            embedding_function=embed,collection_metadata={'hnsw:space':'cosine'})
        try:
            self.store.add_documents(self.chunks,ids=[str(uuid4()) for _ in self.chunks])
        except Exception:
            self.store.delete_collection()
            raise

    def search(self, query: str, k: int = 5) -> list[dict]:
        matches = self.store.similarity_search_with_score(query[:2000],k=min(max(k,1),6))
        return [dict(d.metadata,text=d.page_content,similarity=round(1-float(distance),4))
                for d,distance in matches if float(distance) < .85]

    def close(self):
        self.store.delete_collection()
