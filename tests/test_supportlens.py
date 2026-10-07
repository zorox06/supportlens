from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_core.messages import AIMessage,ToolMessage
from streamlit.testing.v1 import AppTest

from src.agent import EvidenceAssessment,SupportAgent,Understanding,citations
from src.knowledge import SupportKnowledge,load_document
from src.tools import build_tools,repository_name


class TinyEmbeddings(Embeddings):
    """Deterministic test vectors; production always uses Gemini embeddings."""
    def embed_documents(self,texts): return [self.embed_query(t) for t in texts]
    def embed_query(self,text):
        text=text.lower()
        return [float('e104' in text),float('password' in text),.05]


def test_chroma_retrieves_matching_document():
    kb=SupportKnowledge([
        Document(page_content='E104 is a connection timeout. Run Test connection.',metadata={'source':'connection.md'}),
        Document(page_content='Password recovery starts on the sign-in page.',metadata={'source':'account.md'})],
        '',embeddings=TinyEmbeddings())
    try:
        found=kb.search('E104 timeout',1)
        assert found[0]['source']=='connection.md'
        assert found[0]['similarity']>.9
    finally:
        kb.close()


def test_docx_tables_are_ingested():
    word=WordDocument(); table=word.add_table(rows=1,cols=2)
    table.cell(0,0).text='E104'; table.cell(0,1).text='Connection timeout'
    stream=BytesIO(); word.save(stream)
    assert 'E104 | Connection timeout' in load_document('guide.docx',stream.getvalue())[0].page_content


@pytest.mark.parametrize('name,data',[('empty.txt',b''),('script.py',b'print(1)'),('punctuation.pdf',b'not pdf')])
def test_invalid_documents_are_rejected(name,data):
    with pytest.raises(Exception): load_document(name,data)


@pytest.mark.parametrize('repo',['https://evil.example/a/b','a/../b','http://github.com/a/b'])
def test_known_issues_repository_is_constrained(repo):
    with pytest.raises(ValueError): repository_name(repo)


def test_ticket_draft_never_submits():
    ticket=build_tools(None)[2].invoke({'summary':'E104 unresolved','category':'connectivity',
                                       'attempted_steps':['Test connection'],'missing_information':['OS version']})
    assert ticket['status']=='draft' and not ticket['submitted']


def test_grouped_unknown_citation_is_rejected():
    assert not citations('Claim [D-12345678, D-deadbeef]',[{'id':'D-12345678'}])['passed']


SOURCE={'id':'D-12345678','text':'E104 means connection timeout. Run Test connection.',
        'source':'guide.md','chunk':1,'similarity':.9}


class FakeKB:
    def __init__(self,sources): self.sources=sources; self.queries=[]
    def search(self,query,k=5): self.queries.append(query); return self.sources


class FakeModel:
    def __init__(self,replies): self.replies=iter(replies); self.received=[]
    def with_structured_output(self,schema,method=None): return self
    def bind_tools(self,tools): return self
    def invoke(self,messages):
        self.received.append(messages)
        return next(self.replies)


def understanding():
    return Understanding(category='connectivity',search_query='E104 connection timeout',
                         brief='Connection timeout',check_known_issues=False)


def test_supported_request_uses_assessment_and_tool_loop():
    kb=FakeKB([SOURCE])
    model=FakeModel([
        understanding(),EvidenceAssessment(sufficient=True,relevant_ids=[SOURCE['id']],reason='Explicit steps',missing_information=[]),
        AIMessage(content='',tool_calls=[{'name':'search_manuals','args':{'query':'E104'},'id':'search1'}]),
        AIMessage(content='Run Test connection [D-12345678].')])
    result=SupportAgent(kb,'','test',model=model).run('E104 help')
    assert result['checks']['passed']
    assert result['assessment']['sufficient']
    assert len(kb.queries)==2
    assert any(isinstance(m,ToolMessage) for m in model.received[-1])


def test_insufficient_evidence_retries_then_escalates():
    kb=FakeKB([])
    insufficient=EvidenceAssessment(sufficient=False,relevant_ids=[],reason='No coverage',missing_information=['Product version'])
    model=FakeModel([understanding(),insufficient,AIMessage(content='printer paper jam'),insufficient,
                     AIMessage(content='The manuals do not document a printer fix. What product version are you using? Ticket is a draft.')])
    result=SupportAgent(kb,'','test',model=model).run('Printer paper jam')
    assert result['retries']==1
    assert len(kb.queries)==2
    assert not result['ticket']['submitted']
    assert any(e['node']=='escalate' for e in result['trace'])
    assert result['checks']['passed']


def test_hallucinated_assessment_ids_cannot_authorize_a_solution():
    kb=FakeKB([])
    bad=EvidenceAssessment(sufficient=True,relevant_ids=['D-deadbeef'],reason='Claimed',missing_information=[])
    model=FakeModel([understanding(),bad,AIMessage(content='refined query'),bad,AIMessage(content='No supported solution.')])
    result=SupportAgent(kb,'','test',model=model).run('Unrelated problem')
    assert not result['assessment']['sufficient']
    assert 'ticket' in result


def test_interface_loads_before_indexing():
    app=AppTest.from_file(Path(__file__).resolve().parent.parent/'app.py',default_timeout=30).run()
    assert not app.exception
    app.button(key='run').click().run()
    assert not app.exception
    assert 'Index the current documents' in app.error[0].value
